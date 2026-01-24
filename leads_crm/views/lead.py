from django.db import transaction
from django.db.models import Q
from rest_framework import viewsets, status, filters
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.exceptions import PermissionDenied, ValidationError
from django_filters.rest_framework import DjangoFilterBackend
import openpyxl

from accounts.models import User
from leads_crm.models import Lead
from leads_crm.serializers import LeadSerializer


class LeadViewSet(viewsets.ModelViewSet):
    queryset = Lead.objects.all()
    serializer_class = LeadSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [
        DjangoFilterBackend,
        filters.SearchFilter,
        filters.OrderingFilter,
    ]
    filterset_fields = ["stage", "organization"]
    search_fields = ["name", "company_name", "email", "phone", "notes"]
    ordering_fields = [
        "created_at",
        "updated_at",
        "probability",
        "expected_revenue",
        "name",
    ]
    ordering = ["-created_at"]

    def get_queryset(self):
        user = self.request.user
        organization = getattr(user, "organization", None)
        if not organization:
            return Lead.objects.none()

        qs = (
            Lead.objects.select_related(
                "organization", "created_by", "assigned_to", "converted_to_project"
            )
            .prefetch_related("assigned_users")
            .filter(organization=organization)
        )

        if user.is_ceo or user.is_bde:
            return qs

        return qs.filter(Q(assigned_to=user) | Q(assigned_users=user)).distinct()

    def _ensure_can_modify(self, request):
        if not (request.user.is_ceo or request.user.is_bde):
            raise PermissionDenied("You don't have permission to modify leads.")

    def perform_create(self, serializer):
        self._ensure_can_modify(self.request)
        serializer.save()

    def perform_update(self, serializer):
        self._ensure_can_modify(self.request)
        serializer.save()

    def perform_destroy(self, instance):
        self._ensure_can_modify(self.request)
        return super().perform_destroy(instance)

    @action(detail=True, methods=["post"], url_path="assign")
    def assign(self, request, pk=None):
        """Assign one or more users to a lead."""
        self._ensure_can_modify(request)
        lead = self.get_object()
        user_ids = request.data.get("user_ids") or request.data.get("user_id")

        if not user_ids:
            return Response(
                {"error": "user_ids is required."}, status=status.HTTP_400_BAD_REQUEST
            )

        if isinstance(user_ids, str):
            user_ids = [uid.strip() for uid in user_ids.split(",") if uid.strip()]

        try:
            ids = [int(uid) for uid in user_ids]
        except ValueError:
            return Response(
                {"error": "user_ids must be integers."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        valid_users = User.objects.filter(
            id__in=ids,
            organization=lead.organization,
            is_active=True,
        )

        if not valid_users.exists():
            return Response(
                {"error": "No matching active users found for this organization."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        lead.assigned_users.add(*valid_users)
        if not lead.assigned_to:
            lead.assigned_to = valid_users.first()
            lead.save(update_fields=["assigned_to"])

        serializer = self.get_serializer(lead)
        return Response(serializer.data, status=status.HTTP_200_OK)

    @action(detail=False, methods=["post"], url_path="import")
    def import_excel(self, request):
        """Import leads from an Excel file (xlsx)."""
        self._ensure_can_modify(request)
        upload = request.FILES.get("file")
        if not upload:
            return Response(
                {"error": "No file uploaded. Please attach an Excel file."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not upload.name.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
            return Response(
                {"error": "Unsupported file type. Please upload an .xlsx file."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            workbook = openpyxl.load_workbook(upload, data_only=True)
        except Exception as exc:  # noqa: BLE001
            return Response(
                {"error": f"Unable to read Excel file: {exc}"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        normalized = [str(h or "").strip().lower() for h in headers]

        required = ["contact name", "company name", "email", "stage"]
        missing = [col for col in required if col not in normalized]
        if missing:
            return Response(
                {"error": f'Missing required columns: {", ".join(missing)}'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        columns_to_capture = required + ["phone"]
        idx = {
            name: normalized.index(name)
            for name in columns_to_capture
            if name in normalized
        }
        stage_choices = {choice[0] for choice in Lead.STAGE_CHOICES}

        created, errors = 0, []
        user = request.user
        organization = getattr(user, "organization", None)
        if not organization:
            raise ValidationError(
                {"organization": "User is not associated with any organization."}
            )

        with transaction.atomic():
            for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
                values = [cell.value for cell in row]
                try:
                    name = str(values[idx["contact name"]] or "").strip()
                    company_name = str(values[idx["company name"]] or "").strip()
                    email = str(values[idx["email"]] or "").strip()
                    stage = str(values[idx["stage"]] or "NEW").strip().upper()
                    phone = (
                        str(values[idx["phone"]] or "").strip()
                        if "phone" in idx
                        else ""
                    )

                    if not all([name, company_name, email, stage]):
                        raise ValueError(
                            "contact name, company name, email, and stage are required"
                        )

                    if stage not in stage_choices:
                        raise ValueError(f'Invalid stage "{stage}"')

                    lead = Lead.objects.create(
                        organization=organization,
                        name=name,
                        company_name=company_name,
                        email=email,
                        phone=phone or None,
                        stage=stage,
                        created_by=user,
                        assigned_to=user,
                    )
                    lead.assigned_users.add(user)
                    created += 1
                except Exception as exc:  # noqa: BLE001
                    errors.append({"row": row_number, "error": str(exc)})

        return Response(
            {"created": created, "errors": errors}, status=status.HTTP_200_OK
        )
