from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db import transaction
from django.db.models import Q, Count, Sum
from django.http import JsonResponse
from django.utils import timezone
from django.core.paginator import Paginator
import openpyxl

from leads_crm.models import Lead, Proposal, MeetingLog
from accounts.models import User
from projects.models import Project
from django.shortcuts import render


@login_required
def list_leads(request):
    """List all leads with filtering"""
    user = request.user

    # Get user's organization
    if not user.organization:
        messages.error(request, "You are not associated with an organization.")
        return redirect("accounts:dashboard")

    organization = user.organization

    # Base queryset
    leads = (
        Lead.objects.filter(organization=organization)
        .select_related("assigned_to", "created_by", "converted_to_project")
        .prefetch_related("assigned_users")
        .order_by("-created_at")
    )

    if not (user.is_ceo or user.is_bde):
        # Others see only assigned leads
        leads = leads.filter(Q(assigned_to=user) | Q(assigned_users=user))

    # Filtering
    search_query = request.GET.get("search", "")
    stage_filter = request.GET.get("stage", "")
    assigned_filter = request.GET.get("assigned", "")

    if search_query:
        leads = leads.filter(
            Q(name__icontains=search_query)
            | Q(company_name__icontains=search_query)
            | Q(email__icontains=search_query)
            | Q(phone__icontains=search_query)
            | Q(notes__icontains=search_query)
        )

    if stage_filter:
        leads = leads.filter(stage=stage_filter)

    if assigned_filter and (user.is_ceo or user.is_bde):
        leads = leads.filter(assigned_users__id=assigned_filter)

    # Pagination
    paginator = Paginator(leads, 50)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    # Get users for filter dropdown (only for CEO/BDE)
    users = None
    if user.is_ceo or user.is_bde:
        users = User.objects.filter(organization=organization, is_active=True).order_by(
            "first_name", "last_name"
        )

    # Statistics
    total_leads = leads.count()
    new_count = leads.filter(stage="NEW").count()
    contacted_count = leads.filter(stage="CONTACTED").count()
    proposal_sent_count = leads.filter(stage="PROPOSAL_SENT").count()
    negotiation_count = leads.filter(stage="NEGOTIATION").count()
    won_count = leads.filter(stage="CLOSED_WON").count()
    lost_count = leads.filter(stage="CLOSED_LOST").count()
    total_revenue = (
        leads.filter(stage="CLOSED_WON").aggregate(Sum("expected_revenue"))[
            "expected_revenue__sum"
        ]
        or 0
    )
    pipeline_revenue = (
        leads.filter(
            stage__in=["NEW", "CONTACTED", "PROPOSAL_SENT", "NEGOTIATION"]
        ).aggregate(Sum("expected_revenue"))["expected_revenue__sum"]
        or 0
    )

    context = {
        "page_obj": page_obj,
        "leads": page_obj,
        "users": users,
        "total_leads": total_leads,
        "new_count": new_count,
        "contacted_count": contacted_count,
        "proposal_sent_count": proposal_sent_count,
        "negotiation_count": negotiation_count,
        "won_count": won_count,
        "lost_count": lost_count,
        "total_revenue": total_revenue,
        "pipeline_revenue": pipeline_revenue,
        "search_query": search_query,
        "stage_filter": stage_filter,
        "assigned_filter": assigned_filter,
        "can_edit": user.is_ceo or user.is_bde,
    }

    return render(request, "leads_crm/list.html", context)


@login_required
def create_lead(request):
    """Create a new lead"""
    user = request.user

    if not user.organization:
        messages.error(request, "You are not associated with an organization.")
        return redirect("accounts:dashboard")

    # Check permissions
    if not (user.is_ceo or user.is_bde):
        messages.error(request, "You do not have permission to create leads.")
        return redirect("leads_crm:list")

    if request.method == "POST":
        from leads_crm.forms import LeadForm

        form = LeadForm(request.POST, user=user, organization=user.organization)
        if form.is_valid():
            lead = form.save(commit=False)
            lead.organization = user.organization
            lead.created_by = user
            lead.save()
            form.save_m2m()

            if not lead.assigned_users.exists():
                lead.assigned_users.add(user)

            lead.assigned_to = lead.assigned_users.first()
            lead.save(update_fields=["assigned_to"])
            messages.success(request, f"Lead created successfully for {lead.name}!")
            return redirect("leads_crm:list")
    else:
        from leads_crm.forms import LeadForm

        form = LeadForm(user=user, organization=user.organization)

    return render(
        request, "leads_crm/form.html", {"form": form, "title": "Create New Lead"}
    )


@login_required
def update_lead(request, id):
    """Update an existing lead"""
    user = request.user
    lead = get_object_or_404(Lead, id=id)

    # Check permissions
    if not (user.is_ceo or user.is_bde):
        if lead.assigned_to != user:
            messages.error(request, "You do not have permission to edit this lead.")
            return redirect("leads_crm:list")

    if request.method == "POST":
        from leads_crm.forms import LeadForm

        form = LeadForm(
            request.POST, instance=lead, user=user, organization=user.organization
        )
        if form.is_valid():
            lead = form.save(commit=False)
            lead.save()
            form.save_m2m()

            lead.assigned_to = (
                lead.assigned_users.first() if lead.assigned_users.exists() else None
            )
            lead.save(update_fields=["assigned_to"])
            messages.success(request, "Lead updated successfully!")
            return redirect("leads_crm:list")
    else:
        from leads_crm.forms import LeadForm

        form = LeadForm(instance=lead, user=user, organization=user.organization)

    return render(
        request,
        "leads_crm/form.html",
        {"form": form, "lead": lead, "title": "Update Lead"},
    )


@login_required
def convert_to_project(request, id):
    """Convert lead to project"""
    user = request.user

    # Check permissions
    if not user.is_ceo:
        messages.error(request, "Only CEO can convert leads to projects.")
        return redirect("leads_crm:list")

    lead = get_object_or_404(Lead, id=id)

    if lead.converted_to_project:
        messages.warning(request, "This lead has already been converted to a project.")
        return redirect("leads_crm:list")

    if request.method == "POST":
        # Create project from lead
        project = Project.objects.create(
            organization=lead.organization,
            name=f"{lead.company_name} - {lead.name}",
            description=lead.notes or f"Project converted from lead: {lead.name}",
            client_name=lead.name,
            client_company=lead.company_name,
            client_email=lead.email,
            client_phone=lead.phone or "",
            status="NOT_STARTED",
            budget=lead.expected_revenue or 0,
        )

        lead.converted_to_project = project
        lead.stage = "CLOSED_WON"
        lead.save()

        messages.success(request, f"Lead converted to project successfully!")
        return redirect("projects:detail", project.id)

    return render(request, "leads_crm/convert_confirm.html", {"lead": lead})


@login_required
def import_leads(request):
    """Import leads from an Excel file (web UI wrapper)."""
    user = request.user

    if not user.organization:
        messages.error(request, "You are not associated with an organization.")
        return redirect("accounts:dashboard")

    if not (user.is_ceo or user.is_bde):
        messages.error(request, "You do not have permission to import leads.")
        return redirect("leads_crm:list")

    if request.method != "POST":
        return redirect("leads_crm:list")

    upload = request.FILES.get("file")
    if not upload:
        messages.error(request, "Please choose an Excel file to import.")
        return redirect("leads_crm:list")

    if not upload.name.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        messages.error(request, "Unsupported file type. Please upload an .xlsx file.")
        return redirect("leads_crm:list")

    try:
        workbook = openpyxl.load_workbook(upload, data_only=True)
    except Exception as exc:  # noqa: BLE001
        messages.error(request, f"Unable to read Excel file: {exc}")
        return redirect("leads_crm:list")

    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    normalized = [str(h or "").strip().lower() for h in headers]

    required = ["contact name", "company name", "email", "stage"]
    missing = [col for col in required if col not in normalized]
    if missing:
        messages.error(request, f'Missing required columns: {", ".join(missing)}')
        return redirect("leads_crm:list")

    columns_to_capture = required + ["phone"]
    idx = {
        name: normalized.index(name)
        for name in columns_to_capture
        if name in normalized
    }
    stage_choices = {choice[0] for choice in Lead.STAGE_CHOICES}

    created = 0
    errors = []

    with transaction.atomic():
        for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            values = [cell.value for cell in row]

            # Skip completely blank rows
            if not any(values):
                continue

            try:
                name = str(values[idx["contact name"]] or "").strip()
                company_name = str(values[idx["company name"]] or "").strip()
                email = str(values[idx["email"]] or "").strip()
                stage = str(values[idx["stage"]] or "NEW").strip().upper()
                phone = (
                    str(values[idx["phone"]] or "").strip() if "phone" in idx else ""
                )

                if not all([name, company_name, email, stage]):
                    raise ValueError(
                        "contact name, company name, email, and stage are required"
                    )

                if stage not in stage_choices:
                    raise ValueError(f'Invalid stage "{stage}"')

                lead = Lead.objects.create(
                    organization=user.organization,
                    name=name,
                    company_name=company_name,
                    email=email,
                    phone=phone or None,
                    stage=stage,
                    created_by=user,
                    assigned_to=user,
                )
                lead.assigned_users.add(user)
                created += 1
            except Exception as exc:  # noqa: BLE001
                errors.append({"row": row_number, "error": str(exc)})

    if created:
        messages.success(request, f"Imported {created} lead(s) successfully.")
    if errors:
        preview = "; ".join([f"Row {e['row']}: {e['error']}" for e in errors[:3]])
        more = f" (+{len(errors)-3} more)" if len(errors) > 3 else ""
        messages.warning(request, f"Issues on {len(errors)} row(s): {preview}{more}")

    return redirect("leads_crm:list")


@login_required
def list_proposals(request):
    """List proposals (web view)"""
    user = request.user
    if not user.organization:
        messages.error(request, "You are not associated with an organization.")
        return redirect("accounts:dashboard")

    proposals = (
        Proposal.objects.filter(lead__organization=user.organization)
        .select_related("lead", "sent_by")
        .order_by("-sent_at")
    )

    context = {
        "proposals": proposals,
    }
    return render(request, "leads_crm/proposals.html", context)


@login_required
def list_meetings(request):
    """List meeting logs (web view)"""
    user = request.user
    if not user.organization:
        messages.error(request, "You are not associated with an organization.")
        return redirect("accounts:dashboard")

    meetings = (
        MeetingLog.objects.filter(lead__organization=user.organization)
        .select_related("lead", "created_by")
        .order_by("-date")
    )

    context = {
        "meetings": meetings,
    }
    return render(request, "leads_crm/meetings.html", context)
